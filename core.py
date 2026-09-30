"""Shared UI and logic for the Model Armor demo.

The two entry points (cloudrun_app.py and streamlit_app.py) differ only in how
they obtain Google Cloud credentials; everything else lives here.
"""

import hashlib
import io
import os
import re
from dataclasses import dataclass
from typing import Callable, Optional

import streamlit as st
from anthropic import AnthropicVertex
from docx import Document
from google import genai
from google.auth.credentials import Credentials
from google.cloud import modelarmor_v1
from openai import OpenAI
from pypdf import PdfReader

CLOUD_PLATFORM_SCOPE = "https://www.googleapis.com/auth/cloud-platform"

MODEL_OPTIONS = [
    {"name": "gemini-3.8-flash", "display_name": "Gemini 3.8 Flash", "provider": "Google", "location": "global"},
    {"name": "claude-sonnet-5-5", "display_name": "Claude Sonnet 5.5", "provider": "Anthropic", "location": "global"},
    {"name": "gpt-5.6-luna", "display_name": "GPT-5.6 Luna", "provider": "OpenAI"},
]
CLAUDE_MAX_TOKENS = 1024

# Model Armor templates are regional; sanitize calls must target the template's region.
# Defaults to us-central1, which supports every filter this demo offers (asia-southeast1,
# for example, lacks malicious URL, CSAM and multi-language detection). Override with
# GOOGLE_CLOUD_LOCATION, e.g. "us-east1" or the "us"/"eu" multi-regions.
DEFAULT_MODEL_ARMOR_LOCATION = "us-central1"
_LOCATION_PATTERN = re.compile(r"[a-z]+(-[a-z]+[0-9]+)?")

# Detection type -> (template ID or prefix, whether a confidence level is appended).
# Ensure these templates exist in the selected Model Armor location, e.g. "ma-pijb-high".
DETECTION_TYPES = {
    "Malicious URLs": ("ma-mal-url", False),
    "Sensitive data protection (inspect)": ("ma-sdp-inspect", False),
    "Sensitive data protection (de-identify)": ("ma-sdp-deid", False),
    "Prompt injection and jailbreak": ("ma-pijb", True),
    "Responsible AI": ("ma-rai", True),
    "All of the above": ("ma-all", True),
}
CONFIDENCE_LEVELS = {"High only": "high", "Medium and above": "med", "Low and above": "low"}
DEID_TEMPLATE_ID = "ma-sdp-deid"
RESPONSE_TEMPLATE_ID = "ma-all-low"

MAX_UPLOAD_MB = 10
_BYTE_ITEM_TYPE = modelarmor_v1.ByteDataItem.ByteItemType
FILE_TYPES = {
    "pdf": _BYTE_ITEM_TYPE.PDF,
    "docx": _BYTE_ITEM_TYPE.WORD_DOCUMENT,
    "csv": _BYTE_ITEM_TYPE.CSV,
    "txt": _BYTE_ITEM_TYPE.TXT,
}

_MATCH = modelarmor_v1.FilterMatchState.MATCH_FOUND
_NO_MATCH = modelarmor_v1.FilterMatchState.NO_MATCH_FOUND
_SKIPPED = modelarmor_v1.FilterExecutionState.EXECUTION_SKIPPED

RAI_FILTER_TYPES = {
    "sexually_explicit": "Sexually Explicit",
    "hate_speech": "Hate Speech",
    "harassment": "Harassment",
    "dangerous": "Dangerous",
}


@dataclass(frozen=True)
class GoogleAuth:
    """Google Cloud credentials and project, as provided by an entry point."""

    credentials: Optional[Credentials]
    project_id: str
    cache_key: str  # changes whenever the credentials change, so clients are rebuilt
    problem: Optional[str] = None  # why Google Cloud calls can't be made, if they can't


@dataclass(frozen=True)
class Attachment:
    name: str
    ext: str
    data: bytes
    text: str


# ---------------------------------------------------------------------------
# Clients: built lazily per session and rebuilt only when their config changes
# ---------------------------------------------------------------------------

def _cached_client(slot: str, key: tuple, factory: Callable):
    clients = st.session_state.setdefault("clients", {})
    cached = clients.get(slot)
    if cached is not None and cached[0] == key:
        return cached[1]
    client = factory()
    clients[slot] = (key, client)
    return client


def _gemini_client(auth: GoogleAuth, location: str) -> genai.Client:
    return _cached_client(
        "google",
        (auth.cache_key, auth.project_id, location),
        lambda: genai.Client(enterprise=True, credentials=auth.credentials, project=auth.project_id, location=location),
    )


def _anthropic_client(auth: GoogleAuth, region: str) -> AnthropicVertex:
    return _cached_client(
        "anthropic",
        (auth.cache_key, auth.project_id, region),
        lambda: AnthropicVertex(project_id=auth.project_id, region=region, credentials=auth.credentials),
    )


def _openai_client(api_key: str) -> OpenAI:
    key_hash = hashlib.sha256(api_key.encode()).hexdigest()
    return _cached_client("openai", (key_hash,), lambda: OpenAI(api_key=api_key))


def _model_armor_client(auth: GoogleAuth, location: str) -> modelarmor_v1.ModelArmorClient:
    return _cached_client(
        "model_armor",
        (auth.cache_key, location),
        lambda: modelarmor_v1.ModelArmorClient(
            credentials=auth.credentials,
            transport="rest",
            client_options={"api_endpoint": f"modelarmor.{location}.rep.googleapis.com"},
        ),
    )


# ---------------------------------------------------------------------------
# Sidebar helpers
# ---------------------------------------------------------------------------

def _openai_api_key() -> str:
    # Never send a server-side key to the browser: if one is configured, use it
    # without rendering an input for it.
    env_key = os.getenv("OPENAI_API_KEY", "")
    if env_key:
        st.caption("Using the server-configured OpenAI API key.")
        return env_key
    return st.text_input("**OpenAI API key**", type="password")


def _model_armor_location() -> str:
    location = os.getenv("GOOGLE_CLOUD_LOCATION", "").strip() or DEFAULT_MODEL_ARMOR_LOCATION
    # The location becomes part of the endpoint hostname, so accept only location-shaped values.
    if not _LOCATION_PATTERN.fullmatch(location):
        st.error(f"Invalid GOOGLE_CLOUD_LOCATION {location!r}; using {DEFAULT_MODEL_ARMOR_LOCATION}.")
        location = DEFAULT_MODEL_ARMOR_LOCATION
    return location


def _template_id(detection_type: str, confidence_level: Optional[str]) -> str:
    template, needs_confidence = DETECTION_TYPES[detection_type]
    return f"{template}-{CONFIDENCE_LEVELS[confidence_level]}" if needs_confidence else template


# ---------------------------------------------------------------------------
# Attachments
# ---------------------------------------------------------------------------

def _read_attachment(uploaded_file) -> Attachment:
    name = uploaded_file.name
    ext = name.rsplit(".", 1)[-1].lower() if "." in name else ""
    data = uploaded_file.getvalue()

    if ext in ("txt", "csv"):
        text = data.decode("utf-8", errors="replace")
    elif ext == "pdf":
        reader = PdfReader(io.BytesIO(data))
        text = "\n".join(page_text for page in reader.pages if (page_text := page.extract_text()))
    elif ext == "docx":
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append("\t".join(cell.text for cell in row.cells))
        text = "\n".join(parts)
    else:
        raise ValueError(f"Unsupported file type: .{ext}")

    return Attachment(name=name, ext=ext, data=data, text=text)


# ---------------------------------------------------------------------------
# Model Armor results
# ---------------------------------------------------------------------------

def _state_label(result) -> str:
    if result is None:
        return "*Not Assessed*"
    if getattr(result, "execution_state", None) == _SKIPPED:
        return "*Skipped*"
    if result.match_state == _MATCH:
        return "Match Found 🚨"
    if result.match_state == _NO_MATCH:
        return "No Match Found ✅"
    return "*Not Assessed*"


def _filter_results(sanitization_result) -> dict:
    """Map each filter's oneof field name (e.g. 'rai_filter_result') to its result."""
    found = {}
    for filter_result in sanitization_result.filter_results.values():
        kind = modelarmor_v1.FilterResult.pb(filter_result).WhichOneof("filter_result")
        if kind:
            found[kind] = getattr(filter_result, kind)
    return found


def _summarise(sanitization_result) -> tuple[str, Optional[str]]:
    """Return (markdown summary of all filters, de-identified text if any)."""
    found = _filter_results(sanitization_result)

    sdp_result, deid_text = None, None
    sdp = found.get("sdp_filter_result")
    if sdp is not None:
        sdp_kind = modelarmor_v1.SdpFilterResult.pb(sdp).WhichOneof("result")
        if sdp_kind:
            sdp_result = getattr(sdp, sdp_kind)
            if sdp_kind == "deidentify_result" and sdp_result.match_state == _MATCH:
                deid_text = sdp_result.data.text or None

    lines = [
        f"- **Sensitive Data Protection**: {_state_label(sdp_result)}",
        f"- **Prompt Injection and Jailbreak**: {_state_label(found.get('pi_and_jailbreak_filter_result'))}",
        f"- **Malicious URIs**: {_state_label(found.get('malicious_uri_filter_result'))}",
        f"- **CSAM**: {_state_label(found.get('csam_filter_filter_result'))}",
    ]
    if "virus_scan_filter_result" in found:
        lines.append(f"- **Virus Scan**: {_state_label(found['virus_scan_filter_result'])}")

    rai = found.get("rai_filter_result")
    lines.append(f"- **Responsible AI**: {_state_label(rai)}")
    rai_types = rai.rai_filter_type_results if rai is not None else {}
    for key, label in RAI_FILTER_TYPES.items():
        lines.append(f"    - **{label}**: {_state_label(rai_types.get(key))}")

    return "\n".join(lines), deid_text


def _is_match(response) -> bool:
    return response.sanitization_result.filter_match_state == _MATCH


def _sanitize_prompt(client, template_path: str, item: modelarmor_v1.DataItem):
    request = modelarmor_v1.SanitizeUserPromptRequest(name=template_path, user_prompt_data=item)
    return client.sanitize_user_prompt(request=request)


def _sanitize_response(client, template_path: str, text: str):
    request = modelarmor_v1.SanitizeModelResponseRequest(
        name=template_path,
        model_response_data=modelarmor_v1.DataItem(text=text),
    )
    return client.sanitize_model_response(request=request)


def _quote(text: str) -> str:
    return "\n".join(f"> {line}  " for line in text.splitlines())  # trailing spaces keep line breaks


def _render_findings(heading: str, flagged: list) -> str:
    """Render flagged scans inside the current container and return a markdown record."""
    sections = [heading]
    for label, response in flagged:
        summary, deid_text = _summarise(response.sanitization_result)
        sections.append(f"**{label}**\n\n{summary}")
        if deid_text:
            sections.append(f"**De-identified prompt** (not sent to the model)\n\n{_quote(deid_text)}")
    record = "\n\n".join(sections)
    st.markdown(record)
    for label, response in flagged:
        with st.expander(f"Raw result: {label}", expanded=False):
            with st.container(height=300, border=True):
                st.write(response)
    return record


# ---------------------------------------------------------------------------
# Model calls
# ---------------------------------------------------------------------------

def _generate(model: dict, auth: GoogleAuth, openai_key: str, prompt: str) -> str:
    provider = model["provider"]

    if provider == "Google":
        response = _gemini_client(auth, model["location"]).models.generate_content(model=model["name"], contents=prompt)
        if response.text:
            return response.text
        reason = None
        if response.prompt_feedback and response.prompt_feedback.block_reason:
            reason = response.prompt_feedback.block_reason
        elif response.candidates:
            reason = response.candidates[0].finish_reason
        return f"_The model returned no text (reason: {reason or 'unknown'})._"

    if provider == "Anthropic":
        response = _anthropic_client(auth, model["location"]).messages.create(
            model=model["name"],
            max_tokens=CLAUDE_MAX_TOKENS,
            messages=[{"role": "user", "content": prompt}],
        )
        text = "".join(block.text for block in response.content if block.type == "text")
        return text or f"_The model returned no text (stop reason: {response.stop_reason})._"

    response = _openai_client(openai_key).chat.completions.create(
        model=model["name"],
        messages=[{"role": "user", "content": prompt}],
    )
    message = response.choices[0].message
    if message.content:
        return message.content
    if message.refusal:
        return f"_The model refused: {message.refusal}_"
    return f"_The model returned no text (finish reason: {response.choices[0].finish_reason})._"


# ---------------------------------------------------------------------------
# App
# ---------------------------------------------------------------------------

MODEL_ARMOR_ROLE = "model_armor"
FLAGGED_RESPONSE_BANNER = "⚠️ *Flagged by Model Armor — shown here for demonstration. See the findings below.*"


def _chat_message(role: str):
    # Model Armor verdicts get their own bubble so they are never mistaken for the model's output.
    if role == MODEL_ARMOR_ROLE:
        return st.chat_message("Model Armor", avatar="🛡️")
    return st.chat_message(role)


def run(auth_ui: Callable[[], GoogleAuth]) -> None:
    """Render the app. `auth_ui` draws the credential/project inputs and returns them."""
    st.set_page_config(page_title="Model Armor Demo", page_icon="🛡️", initial_sidebar_state="auto")
    st.session_state.setdefault("messages", [])

    with st.sidebar:
        st.title("🛡️ Model Armor Demo")
        with st.expander("**⚙️ Model Settings**", expanded=False):
            model = st.selectbox("**Model**", options=MODEL_OPTIONS, format_func=lambda m: m["display_name"])
            openai_key = _openai_api_key() if model["provider"] == "OpenAI" else ""

        with st.expander("**⚙️ Model Armor Settings**", expanded=True):
            with st.expander("**Project Settings**", expanded=False):
                auth = auth_ui()
                location = _model_armor_location()
                st.text_input("**Location**", value=location, disabled=True, help="Set by the GOOGLE_CLOUD_LOCATION environment variable.")

            with st.expander("**Detection Settings**", expanded=True):
                template_id = None
                sanitize_request = st.checkbox("Sanitize prompt request?")
                if sanitize_request:
                    detection_type = st.radio("**Detection type**", list(DETECTION_TYPES))
                    confidence_level = None
                    if DETECTION_TYPES[detection_type][1]:
                        confidence_level = st.radio("**Confidence level**", list(CONFIDENCE_LEVELS))
                    template_id = _template_id(detection_type, confidence_level)
                sanitize_response = st.checkbox("Sanitize model response?", help=f"Uses the `{RESPONSE_TEMPLATE_ID}` template")

    for message in st.session_state.messages:
        with _chat_message(message["role"]):
            st.markdown(message["content"])

    prompt = st.chat_input(
        "Ask anything",
        accept_file=True,
        file_type=list(FILE_TYPES),
        max_upload_size=MAX_UPLOAD_MB,
    )
    if not prompt:
        return

    # Validate configuration before doing anything
    needs_google = model["provider"] in ("Google", "Anthropic") or sanitize_request or sanitize_response
    if needs_google and auth.problem:
        st.error(auth.problem)
        st.stop()
    if model["provider"] == "OpenAI" and not openai_key:
        st.error("Please provide the OpenAI API key.")
        st.stop()

    prompt_text = prompt.text or ""
    attachment = None
    if prompt.files:
        try:
            attachment = _read_attachment(prompt.files[0])
        except Exception as e:
            st.error(f"Could not read {prompt.files[0].name}: {e}")
            st.stop()

    full_prompt = prompt_text
    if attachment and attachment.text:
        full_prompt += f"\n\n[File Content]\n{attachment.text}"
    if not full_prompt.strip():
        source = attachment.name if attachment else "your message"
        st.error(f"No text could be extracted from {source}. Add a message or try another file.")
        st.stop()

    display_text = prompt_text + (f"\n\nFile attached: {attachment.name}" if attachment else "")
    with st.chat_message("user"):
        st.markdown(display_text)
    st.session_state.messages.append({"role": "user", "content": display_text})

    armor = _model_armor_client(auth, location) if (sanitize_request or sanitize_response) else None
    template_base = f"projects/{auth.project_id}/locations/{location}/templates"

    # Prompt sanitisation
    if sanitize_request:
        # Files are scanned natively by Model Armor, except for de-identification,
        # which Model Armor only supports on text.
        scans = []
        if attachment and template_id != DEID_TEMPLATE_ID:
            if prompt_text:
                scans.append(("Prompt text", modelarmor_v1.DataItem(text=prompt_text)))
            byte_item = modelarmor_v1.ByteDataItem(byte_data_type=FILE_TYPES[attachment.ext], byte_data=attachment.data)
            scans.append((f"File: {attachment.name}", modelarmor_v1.DataItem(byte_item=byte_item)))
        else:
            scans.append(("Prompt", modelarmor_v1.DataItem(text=full_prompt)))

        try:
            with st.spinner("Analysing prompt request..."):
                results = [
                    (label, _sanitize_prompt(armor, f"{template_base}/{template_id}", item))
                    for label, item in scans
                ]
        except Exception as e:
            st.error(f"Model Armor error during request sanitisation: {e}")
            st.stop()

        flagged = [(label, response) for label, response in results if _is_match(response)]
        if flagged:
            with _chat_message(MODEL_ARMOR_ROLE):
                record = _render_findings(f"🚨 **Prompt blocked** (template `{template_id}`)", flagged)
            st.session_state.messages.append({"role": MODEL_ARMOR_ROLE, "content": record})
            st.stop()

    # Model response
    try:
        with st.spinner("Generating response..."):
            model_response = _generate(model, auth, openai_key, full_prompt)
    except Exception as e:
        error_text = f"Error generating LLM response: {e}"
        with st.chat_message("assistant"):
            st.error(error_text)
        st.session_state.messages.append({"role": "assistant", "content": error_text})
        st.stop()

    # Response sanitisation happens before the response is shown, so a flagged
    # response is labelled as such from the moment it appears.
    scan, scan_error = None, None
    if sanitize_response:
        try:
            with st.spinner("Analysing model response..."):
                scan = _sanitize_response(armor, f"{template_base}/{RESPONSE_TEMPLATE_ID}", model_response)
        except Exception as e:
            scan_error = f"Model Armor error during response sanitisation: {e}"

    response_flagged = scan is not None and _is_match(scan)
    shown_response = f"{FLAGGED_RESPONSE_BANNER}\n\n{model_response}" if response_flagged else model_response
    with st.chat_message("assistant"):
        st.markdown(shown_response)
    st.session_state.messages.append({"role": "assistant", "content": shown_response})

    if scan_error:
        st.error(scan_error)
        st.stop()

    if response_flagged:
        with _chat_message(MODEL_ARMOR_ROLE):
            record = _render_findings(
                f"⚠️ **Model response flagged** (template `{RESPONSE_TEMPLATE_ID}`)",
                [("Model response", scan)],
            )
        st.session_state.messages.append({"role": MODEL_ARMOR_ROLE, "content": record})
